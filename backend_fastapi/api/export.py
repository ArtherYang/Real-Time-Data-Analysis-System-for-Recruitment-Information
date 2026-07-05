"""数据导出 API — FastAPI 版本"""

import csv
import io
from datetime import date
from typing import Optional

from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse

from app.database import get_db
from app.models.job import Job

router = APIRouter(tags=["数据导出"])


@router.get("/export/csv")
def export_csv(
    city: Optional[str] = Query(None),
    job_category: Optional[str] = Query(None),
    platform: Optional[str] = Query(None),
    experience: Optional[str] = Query(None),
    education: Optional[str] = Query(None),
):
    session = next(get_db())
    try:
        q = session.query(Job).filter(Job.status == "有效")
        if city:          q = q.filter(Job.city == city)
        if job_category:  q = q.filter(Job.job_category == job_category)
        if platform:      q = q.filter(Job.platform == platform)
        if experience:    q = q.filter(Job.experience == experience)
        if education:     q = q.filter(Job.education == education)

        jobs = q.all()
        output = io.StringIO()
        output.write("﻿")
        fields = ["title", "company", "city", "salary_min", "salary_max",
                   "salary_type", "experience", "education", "job_category",
                   "skills", "platform", "published_at", "company_type", "company_size", "welfare"]
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for job in jobs:
            writer.writerow(job.to_dict())

        today_str = date.today().isoformat()
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename=recruitment_data_{today_str}.csv"},
        )
    finally:
        session.close()


@router.get("/export/excel")
def export_excel(
    city: Optional[str] = Query(None),
    job_category: Optional[str] = Query(None),
    platform: Optional[str] = Query(None),
    experience: Optional[str] = Query(None),
    education: Optional[str] = Query(None),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return {"code": 500, "message": "请安装 openpyxl: pip install openpyxl"}

    session = next(get_db())
    try:
        q = session.query(Job).filter(Job.status == "有效")
        if city:          q = q.filter(Job.city == city)
        if job_category:  q = q.filter(Job.job_category == job_category)
        if platform:      q = q.filter(Job.platform == platform)
        if experience:    q = q.filter(Job.experience == experience)
        if education:     q = q.filter(Job.education == education)
        jobs = q.all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "招聘数据"

        header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="409EFF", end_color="409EFF", fill_type="solid")
        thin_border = Border(left=Side(style="thin"), right=Side(style="thin"),
                             top=Side(style="thin"), bottom=Side(style="thin"))

        headers = ["岗位名称", "公司", "城市", "最低薪资", "最高薪资", "薪资类型",
                    "经验要求", "学历要求", "岗位分类", "技能标签", "数据来源",
                    "发布日期", "公司类型", "公司规模", "福利"]
        fields = ["title", "company", "city", "salary_min", "salary_max", "salary_type",
                   "experience", "education", "job_category", "skills", "platform",
                   "published_at", "company_type", "company_size", "welfare"]

        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font, cell.fill, cell.alignment = header_font, header_fill, Alignment(horizontal="center")
            cell.border = thin_border

        for ri, job in enumerate(jobs, 2):
            d = job.to_dict()
            for ci, f in enumerate(fields, 1):
                cell = ws.cell(row=ri, column=ci, value=d.get(f, ""))
                cell.border = thin_border

        for ci, w in enumerate([22,18,8,14,14,10,10,10,12,35,12,14,12,14,30], 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A2"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        today_str = date.today().isoformat()
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=recruitment_data_{today_str}.xlsx"},
        )
    finally:
        session.close()
